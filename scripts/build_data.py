#!/usr/bin/env python3
"""
Đọc Google Sheet (tải về dạng .xlsx), kiểm tra, rồi gộp vào data.json cho web.

- Tuần trong tab "Số liệu tuần" được ghi đè vào đúng tuần đó; các tuần cũ trong data.json giữ nguyên.
- Thành viên, cáo thị, cập nhật game, hũ thưởng, tài liệu, cài đặt: lấy theo Sheet.
- Dữ liệu sai thì dừng, không ghi gì (web giữ bản cũ).

Dùng:
  python scripts/build_data.py --sheet-id <ID>            # tải từ Google (Sheet chia sẻ "Bất kỳ ai có link")
  python scripts/build_data.py --xlsx path/to/file.xlsx   # đọc file có sẵn (để thử)
Biến môi trường: ADMIN_PASS (mật khẩu quản trị web), FORCE=1 (cho phép ghi đè tuần cũ hơn tuần mới nhất).
"""
import argparse, datetime as dt, hashlib, io, json, os, re, sys, unicodedata, urllib.request, urllib.parse

T_SET, T_MEM, T_WEEK, T_PASTE = 'Cài đặt', 'Thành viên', 'Số liệu tuần', 'Dán ảnh'
T_NOTICE, T_UPD, T_DONOR, T_DOCS = 'Cáo thị', 'Cập nhật game', 'Mạnh thường quân', 'Tàng thư các'
HE_TU = ['Pháp tu', 'Thể tu', 'Nho tu', 'Ngự quỷ', 'Ngự kiếm']
HE_OLD = {'tu pháp': 'Pháp tu', 'luyện thể': 'Thể tu', 'nho thánh': 'Nho tu', 'ngự quỷ': 'Ngự quỷ', 'ngự kiếm': 'Ngự kiếm', 'pháp tu': 'Pháp tu', 'thể tu': 'Thể tu', 'nho tu': 'Nho tu'}
LEFT_STATUS = ('Nghỉ hẳn', 'Bị kick')
VN = dt.timezone(dt.timedelta(hours=7))
ERRORS, WARN = [], []

def err(msg): ERRORS.append(msg)
def warn(msg): WARN.append(msg)
def nfc(s): return unicodedata.normalize('NFC', str(s)).strip() if s is not None else ''
def key(s): return re.sub(r'[\s•·.\-_|]', '', nfc(s).lower())

def num(v, where=''):
    if v is None or v == '': return None
    if isinstance(v, bool): return int(v)
    if isinstance(v, (int, float)): return v
    s = nfc(v).replace(' ', '')
    if re.fullmatch(r'[\d.]+', s) and s.count('.') >= 1 and len(s.split('.')[-1]) == 3: s = s.replace('.', '')  # 10.868 kiểu VN
    s = s.replace(',', '.')
    try: return float(s) if '.' in s else int(s)
    except ValueError:
        err(f'{where}: "{v}" không phải số'); return None

def parse_lc(v, where=''):
    if v is None or v == '': return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return int(v * 1e9) if v < 1000 else int(v)        # gõ 11.09 hiểu là 11.09B
    s = nfc(v).upper().replace(',', '.').replace(' ', '')
    m = re.fullmatch(r'([\d.]+)(K|M|B|TỶ|TY|TR)?', s)
    if not m: err(f'{where}: lực chiến "{v}" không đọc được (ví dụ đúng: 11.09B)'); return None
    k = {'K': 1e3, 'M': 1e6, 'TR': 1e6, 'B': 1e9, 'TỶ': 1e9, 'TY': 1e9}.get(m.group(2) or '', 1)
    x = float(m.group(1)) * k
    return int(x * 1e9) if (not m.group(2) and x < 1000) else int(x)

def tiers_pct(sv):
    # Phong Thần Bảng chia theo % số người, xếp theo lực chiến. Hoàng Bảng là phần còn lại.
    p = [float(num(sv(f'{n} Bảng (% người)', d), 'Cài đặt') or 0) for n, d in (('Thiên', 10), ('Địa', 20), ('Huyền', 30))]
    if sum(p) >= 100: err(f'Cài đặt: tổng % Thiên + Địa + Huyền Bảng = {sum(p):g}%, phải nhỏ hơn 100%')
    p = [int(x) if x == int(x) else x for x in p]
    return [{'name': 'Thiên', 'pct': p[0]}, {'name': 'Địa', 'pct': p[1]}, {'name': 'Huyền', 'pct': p[2]},
            {'name': 'Hoàng', 'pct': round(100 - sum(p), 2)}]

def boolv(v):
    if isinstance(v, bool): return v
    return nfc(v).lower() in ('true', 'x', '1', 'có', 'co', 'yes', '✓', 'v')

def to_date(v):
    if v in (None, ''): return None
    if isinstance(v, dt.datetime): return v
    if isinstance(v, dt.date): return dt.datetime(v.year, v.month, v.day)
    s = nfc(v)
    for f in ('%d/%m/%Y %H:%M', '%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%d/%m/%y'):
        try: return dt.datetime.strptime(s, f)
        except ValueError: pass
    err(f'Ngày "{v}" không đọc được (ví dụ đúng: 25/09/2026)'); return None

def iso(d): return d.strftime('%Y-%m-%d') if d else ''
def iso_week(d): y, w, _ = d.isocalendar(); return f'{y}-W{w:02d}'
def game_week(d):  # tuần game đổi lúc 8h sáng thứ Hai
    return iso_week(d - dt.timedelta(hours=8))
def norm_week(v, fallback_date=None):
    if v in (None, ''):
        return game_week(fallback_date) if fallback_date else None
    s = nfc(v).upper().replace(' ', '')
    m = re.fullmatch(r'(\d{4})-?W(\d{1,2})', s)
    if m: return f'{m.group(1)}-W{int(m.group(2)):02d}'
    m = re.fullmatch(r'(?:T|TUẦN|TUAN|W)?(\d{1,2})', s)
    if m and fallback_date: return f'{fallback_date.isocalendar()[0]}-W{int(m.group(1)):02d}'
    err(f'Mã tuần "{v}" sai, ví dụ đúng: 2026-W41'); return None

# ---------------- đọc workbook ----------------
def load_wb(args):
    import openpyxl
    if args.xlsx:
        return openpyxl.load_workbook(args.xlsx, data_only=True)
    api, k = os.environ.get('SHEET_API', '').strip(), os.environ.get('SHEET_KEY', '').strip()
    if api and k:
        # Sheet riêng tư: Apps Script (chạy bằng quyền chủ Sheet) trả file xlsx khi đúng khoá
        import base64
        url = api + ('&' if '?' in api else '?') + 'key=' + urllib.parse.quote(k)
        with urllib.request.urlopen(urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'}), timeout=120) as r:
            txt = r.read().decode('utf-8', 'replace').strip()
        try: data = base64.b64decode(txt, validate=True)
        except Exception: data = b''
        if not data.startswith(b'PK'):
            sys.exit('Không tải được Sheet qua Apps Script (' + txt[:80] + '). Kiểm tra secret SHEET_API, SHEET_KEY và bản triển khai Web app.')
    else:
        url = f'https://docs.google.com/spreadsheets/d/{args.sheet_id}/export?format=xlsx'
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=60) as r:
            data = r.read()
        if not data.startswith(b'PK'):
            sys.exit('Không tải được Sheet. Kiểm tra Sheet đã chia sẻ "Bất kỳ ai có đường liên kết đều có thể xem" và SHEET_ID đúng, hoặc cài chế độ Sheet riêng tư (SHEET_API, SHEET_KEY).')
    return openpyxl.load_workbook(io.BytesIO(data), data_only=True)

def rows(ws, start=2, ncol=None):
    out = []
    for r in ws.iter_rows(min_row=start, values_only=True):
        r = list(r[:ncol] if ncol else r)
        if ncol and len(r) < ncol: r += [None] * (ncol - len(r))
        if all(v in (None, '') for v in r): continue
        out.append(r)
    return out

def need(wb, name):
    if name not in wb.sheetnames: err(f'Thiếu tab "{name}"'); return None
    return wb[name]

def check_header(ws, row, expect, tab):
    got = [nfc(c.value) for c in ws[row][:len(expect)]]
    if got != expect: err(f'Tab "{tab}" dòng {row}: tiêu đề cột phải là {expect}, đang là {got}')

# ---------------- main ----------------
SEGS = 8   # vòng quay 8 ô xen kẽ: ô chẵn Chia thưởng, ô lẻ Tích trữ
MAX_KEEP = 3   # tích trữ tối đa 3 lần liên tiếp, lần thứ 4 chắc chắn nổ hũ

def spin_key(pw, token):
    """Mã hoá token GitHub (chỉ có quyền chạy Actions) bằng mật khẩu quản trị, để nút Quay trên web gọi được máy quay.
    Không có mật khẩu thì không giải mã được. Muối và IV suy ra cố định để data.json không đổi giữa các lần chạy."""
    if not pw or not token: return None
    import base64
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM
    salt = hashlib.sha256(('salt|' + pw + '|' + token).encode()).digest()[:16]
    iv = hashlib.sha256(('iv|' + pw + '|' + token).encode()).digest()[:12]
    key = hashlib.pbkdf2_hmac('sha256', pw.encode(), salt, 150000, 32)
    ct = AESGCM(key).encrypt(iv, token.encode(), None)
    b = lambda x: base64.b64encode(x).decode()
    return {'salt': b(salt), 'iv': b(iv), 'ct': b(ct), 'iter': 150000}

def qd_of(e):
    v = (e.get('acts') or {}).get('quyetDau', 0)
    return 3 if v is True else max(0, min(3, int(v or 0)))

ELIG_DEF = {'qd': 3, 'bh': 2, 'qn': 2, 'ch': 0, 'dg': 0, 'ld': False}

def eligible_names(weeks, members, wk, c=None):
    """Danh sách đủ điều kiện quay thưởng tuần wk, theo điều kiện ở tab Cài đặt (giống hệt cách web tính)."""
    c = dict(ELIG_DEF, **(c or {}))
    ws = [w for w in sorted(weeks) if w <= wk]
    def missed(mid, key, n):
        if not n or len(ws) < n: return False
        return all((weeks[w].get(mid) is not None) and not (weeks[w][mid].get('acts') or {}).get(key) for w in ws[-n:])
    out = []
    for m in members:
        if m.get('leftAt'): continue
        e = weeks.get(wk, {}).get(m['id'])
        if not e: continue
        a = e.get('acts') or {}
        if qd_of(e) < c['qd']: continue
        if c['ld'] and not a.get('loanDau'): continue
        if c['ch'] and (e.get('ch') or 0) < c['ch']: continue
        if c['dg'] and (e.get('dg') or 0) < c['dg']: continue
        if missed(m['id'], 'batHoang', c['bh']) or missed(m['id'], 'quyNhat', c['qn']): continue
        out.append(m['name'])
    return out

def spin_week(now, wk, weeks, members, hist, donors, elig=None):
    """Quay thưởng cho tuần wk khi quản trị bấm nút trên web. Kết quả do máy bốc ngẫu nhiên (secrets), không ai chọn được:
    50% Chia thưởng (bốc 1 người đủ điều kiện, nhận cả hũ), 50% Tích trữ (hũ giữ nguyên). Mỗi tuần chỉ quay 1 lần."""
    import secrets
    if not wk or wk not in weeks: print(f'Không quay: chưa có số liệu tuần {wk}.'); return None
    if any(h.get('week') == wk for h in hist): print(f'Không quay: tuần {wk} đã quay rồi.'); return None
    names = eligible_names(weeks, members, wk, elig)
    today = now.strftime('%Y-%m-%d')
    sp = {'week': wk, 'date': today, 'at': now.strftime('%H:%M %d/%m/%Y'), 'count': len(names), 'list': names, 'auto': True}
    # Bảo hiểm: đã Tích trữ 3 lần liên tiếp thì lần này chắc chắn Chia thưởng (nếu có người đủ điều kiện)
    streak = 0
    for h in sorted(hist, key=lambda x: (x.get('date', ''), x.get('week', '')), reverse=True):
        if not h.get('keep'): break
        streak += 1
    pity = bool(names) and streak >= MAX_KEEP
    win = pity or (bool(names) and secrets.randbelow(2) == 0)
    if pity: sp['pity'] = True
    sp['seg'] = 2 * secrets.randbelow(SEGS // 2) + (0 if win else 1)
    if win:
        paid = sum(h.get('amount', 0) for h in hist if not h.get('keep'))
        sp['winner'] = names[secrets.randbelow(len(names))]
        sp['amount'] = sum(x['amount'] for x in donors if x['date'] <= today) - paid
    else:
        sp.update({'winner': 'Tích trữ', 'keep': True, 'amount': 0})
        if not names: sp['none'] = True
    return sp

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--sheet-id'); ap.add_argument('--xlsx'); ap.add_argument('--out', default='data.json')
    args = ap.parse_args()
    if not args.sheet_id and not args.xlsx and not os.environ.get('SHEET_API'): sys.exit('Cần --sheet-id, --xlsx hoặc SHEET_API')
    wb = load_wb(args)
    old = {}
    if os.path.exists(args.out):
        with open(args.out, encoding='utf-8') as f: old = json.load(f)
        if old.get('demo'): old = {}   # data.json đang là dữ liệu thử => bỏ, dựng mới hoàn toàn từ Sheet
    now = dt.datetime.now(VN)

    # ---- Cài đặt
    st = {}
    ws = need(wb, T_SET)
    if ws:
        for r in rows(ws, 2, 2):
            if r[0]: st[nfc(r[0])] = r[1]
    def sv(k, d=None): v = st.get(k); return d if v in (None, '') else v
    roles = [nfc(x) for x in str(sv('Chức vụ', 'Tông Chủ, Phó Tông Chủ, Trưởng Lão, Chủ Sự, Thành Viên')).split(',') if nfc(x)]
    pw = os.environ.get('ADMIN_PASS', '')
    settings = {
        'name': nfc(sv('Tên tông', 'Tông Môn')), 'sub': nfc(sv('Dòng phụ', 'Nhất Niệm Tiêu Dao')),
        'maxMembers': int(num(sv('Số thành viên tối đa', 50), 'Cài đặt') or 50),
        'kpi': {'ch': num(sv('KPI cống hiến mỗi tuần', 7000), 'Cài đặt') or 0, 'dg': num(sv('KPI lệnh dị giới mỗi tuần', 0), 'Cài đặt') or 0},
        'tiers': tiers_pct(sv),
        'roles': roles, 'potUnit': nfc(sv('Đơn vị hũ thưởng', 'VNĐ')),
        'noticeMax': int(num(sv('Số cáo thị hiển thị', 4), 'Cài đặt') or 4),
        'elig': {'qd': int(num(sv('Quay thưởng: Quyết đấu tối thiểu (lượt)', 3), 'Cài đặt') or 0),
                 'bh': int(num(sv('Quay thưởng: loại nếu bỏ Bát hoang (tuần liền)', 2), 'Cài đặt') or 0),
                 'qn': int(num(sv('Quay thưởng: loại nếu bỏ Quy nhất (tuần liền)', 2), 'Cài đặt') or 0),
                 'ch': num(sv('Quay thưởng: cống hiến tối thiểu', 0), 'Cài đặt') or 0,
                 'dg': num(sv('Quay thưởng: lệnh dị giới tối thiểu', 0), 'Cài đặt') or 0,
                 'ld': key(nfc(sv('Quay thưởng: bắt buộc Loạn đấu', 'Không'))) in ('có', 'co', 'x', '1', 'true')},
        'sheetUrl': nfc(sv('Link Google Sheet', '')),
        'adminHash': hashlib.sha256(pw.encode()).hexdigest() if pw else (old.get('settings') or {}).get('adminHash'),
        'spinKey': spin_key(pw, os.environ.get('SPIN_TOKEN', '')) or (old.get('settings') or {}).get('spinKey'),
        'repo': os.environ.get('GITHUB_REPOSITORY') or (old.get('settings') or {}).get('repo', ''),
    }

    # ---- Thành viên
    old_members = {m['id']: m for m in old.get('members', [])}
    by_key = {}
    for m in old.get('members', []):
        for n in [m['name']] + m.get('aliases', []): by_key.setdefault(key(n), m['id'])
    members, seen_ids = [], set()
    ws = need(wb, T_MEM)
    if ws:
        check_header(ws, 1, ['Tên ingame', 'Tên Zalo', 'Chức vụ', 'Nhánh chủ tu', 'Nhóm Thông báo', 'Nhóm Trò chuyện', 'Trạng thái', 'Ngày rời', 'Lý do rời', 'Tên cũ', 'Ghi chú'], T_MEM)
        for i, r in enumerate(rows(ws, 2, 11), start=2):
            name = nfc(r[0])
            if not name: continue
            aliases = [nfc(x) for x in str(r[9] or '').split(',') if nfc(x)]
            mid = next((by_key[key(n)] for n in [name] + aliases if key(n) in by_key), None)
            if mid in seen_ids: err(f'Tab Thành viên dòng {i}: "{name}" bị trùng'); continue
            base = dict(old_members.get(mid, {})) if mid else {}
            if not mid:
                mid = 'm' + hashlib.md5(name.encode()).hexdigest()[:8]
                while mid in old_members or mid in seen_ids: mid += 'x'
            seen_ids.add(mid)
            he = nfc(r[3])
            if he: he = HE_OLD.get(he.lower(), he)
            if he and he not in HE_TU: err(f'Tab Thành viên dòng {i}: Nhánh chủ tu "{r[3]}" không hợp lệ (chọn: {", ".join(HE_TU)})')
            stt = nfc(r[6]) or 'Đang ở'
            if stt not in ('Đang ở', 'Tạm nghỉ') + LEFT_STATUS: err(f'Tab Thành viên dòng {i}: Trạng thái "{stt}" không hợp lệ'); stt = 'Đang ở'
            m = {'id': mid, 'name': name, 'zalo': nfc(r[1]), 'role': nfc(r[2]) or 'Thành Viên', 'heChinh': he,
                 'groups': {'thongBao': boolv(r[4]), 'troChuyen': boolv(r[5])}, 'away': stt == 'Tạm nghỉ',
                 'joinedAt': base.get('joinedAt') or iso(now), 'aliases': aliases, 'note': nfc(r[10]), 'lc': base.get('lc', 0)}
            if stt in LEFT_STATUS:
                m['leftAt'] = iso(to_date(r[7])) or base.get('leftAt') or iso(now)
                m['leftReason'] = stt + (': ' + nfc(r[8]) if nfc(r[8]) else '')
                m['blacklist'] = stt == 'Bị kick'
            members.append(m)
            for n in [name] + aliases: by_key[key(n)] = mid
    # người có trong data cũ nhưng đã xoá khỏi Sheet: giữ lại như cựu thành viên để không mất lịch sử
    for mid, m in old_members.items():
        if mid not in seen_ids:
            m = dict(m); m.setdefault('leftAt', iso(now)); m.setdefault('leftReason', 'Không còn trong Sheet')
            members.append(m); seen_ids.add(mid)
    mem_by_id = {m['id']: m for m in members}

    # ---- Số liệu tuần
    weeks = dict(old.get('weeks', {})); sect = dict(old.get('sect', {}))
    ws = need(wb, T_WEEK)
    wk = None
    if ws:
        wk = norm_week(ws['B1'].value)
        check_header(ws, 4, ['Tên ingame', 'Lực chiến', 'Cống hiến', 'Online', 'Lệnh dị giới', 'Quyết đấu (lượt)', 'Loạn đấu', 'Bát hoang', 'Quy nhất'], T_WEEK)
    entries = {}
    if ws and wk:
        latest = max(weeks) if weeks else None
        if latest and wk < latest and os.environ.get('FORCE') != '1':
            err(f'Tuần trong Sheet ({wk}) cũ hơn tuần mới nhất trên web ({latest}). Sửa ô B1, hoặc chạy lại với FORCE=1 nếu cố ý sửa tuần cũ.')
        for i, r in enumerate(rows(ws, 5, 9), start=5):
            name = nfc(r[0])
            if not name: continue
            lc, ch, seen, dg = parse_lc(r[1], f'Số liệu tuần dòng {i}'), num(r[2], f'Số liệu tuần dòng {i} cột Cống hiến'), nfc(r[3]), num(r[4], f'Số liệu tuần dòng {i} cột Lệnh dị giới')
            qd = num(r[5], f'Số liệu tuần dòng {i} cột Quyết đấu') or 0
            if qd not in (0, 1, 2, 3): err(f'Số liệu tuần dòng {i}: Quyết đấu phải từ 0 đến 3'); qd = 0
            acts = {'quyetDau': int(qd), 'loanDau': boolv(r[6]), 'batHoang': boolv(r[7]), 'quyNhat': boolv(r[8])}
            if lc is None and ch is None and dg is None and not seen and not qd and not any(list(acts.values())[1:]): continue
            mid = by_key.get(key(name))
            if not mid:
                warn(f'"{name}" có trong Số liệu tuần nhưng chưa có trong tab Thành viên, đã tự thêm')
                mid = 'm' + hashlib.md5(name.encode()).hexdigest()[:8]
                mm = {'id': mid, 'name': name, 'zalo': '', 'role': 'Thành Viên', 'heChinh': '', 'groups': {}, 'joinedAt': iso(now), 'aliases': [], 'note': '', 'lc': 0}
                members.append(mm); mem_by_id[mid] = mm; by_key[key(name)] = mid
            e = {'acts': acts}
            if lc is not None: e['lc'] = lc
            if ch is not None: e['ch'] = ch
            if dg is not None: e['dg'] = dg
            if seen: e['seen'] = seen
            entries[mid] = e
        bh, qn, kho = num(ws['B2'].value, 'Hạng Bát hoang'), num(ws['D2'].value, 'Hạng Quy nhất'), num(ws['F2'].value, 'Kho lệnh dị giới')
        sc = {k: v for k, v in (('bh', bh), ('qn', qn), ('kho', kho)) if v is not None}
        if entries: weeks[wk] = entries
        if sc: sect[wk] = sc
        if entries and wk >= (max(weeks) if weeks else wk):
            for mid, e in entries.items():
                if e.get('lc'): mem_by_id[mid]['lc'] = e['lc']
    # lực chiến hiện tại = lần gần nhất có số
    for m in members:
        for w in sorted(weeks, reverse=True):
            e = weeks[w].get(m['id'])
            if e and e.get('lc'): m['lc'] = e['lc']; break

    # ---- Cáo thị / Cập nhật game
    notices, updates = [], []
    ws = need(wb, T_NOTICE)
    if ws:
        for i, r in enumerate(rows(ws, 2, 5), start=2):
            if not nfc(r[2]) and not nfc(r[1]): continue
            d = to_date(r[0]) or now.replace(tzinfo=None)
            notices.append({'id': f'n{i}', 'title': nfc(r[1]), 'text': nfc(r[2]), 'by': nfc(r[3]), 'pinned': boolv(r[4]),
                            'at': d.strftime('%d/%m/%Y %H:%M') if (d.hour or d.minute) else d.strftime('%d/%m/%Y'), 'ts': int(d.timestamp())})
    ws = need(wb, T_UPD)
    if ws:
        for i, r in enumerate(rows(ws, 2, 4), start=2):
            if not nfc(r[1]): continue
            d = to_date(r[0])
            if not d: err(f'Tab Cập nhật game dòng {i}: thiếu ngày'); continue
            updates.append({'id': f'u{i}', 'date': iso(d), 'title': nfc(r[1]), 'note': nfc(r[2]), 'pinned': boolv(r[3]), 'ts': i})

    # ---- Hũ thưởng: tự tính = tổng góp - tổng đã trả. Trúng Chia thưởng thì nhận toàn bộ hũ tại thời điểm quay.
    pot = {'value': 0, 'donors': [], 'history': []}
    ws = need(wb, T_DONOR)
    if ws:
        for i, r in enumerate(rows(ws, 2, 4), start=2):
            if not nfc(r[1]): continue
            d = to_date(r[0])
            if not d: err(f'Mạnh thường quân dòng {i}: thiếu ngày góp'); continue
            amt = num(r[2], f'Mạnh thường quân dòng {i}') or 0
            if amt <= 0: err(f'Mạnh thường quân dòng {i}: số tiền phải lớn hơn 0')
            pot['donors'].append({'name': nfc(r[1]), 'amount': amt, 'date': iso(d), 'note': nfc(r[3])})
    # Lịch sử quay do máy tự quay, lưu trong data.json (không nhập trên Sheet nữa)
    hist = [h for h in (old.get('pot') or {}).get('history', []) if isinstance(h, dict)]
    # Quay cho tuần mới nhất đang có trên web (trùng với tuần web hiển thị ở vòng quay)
    sp = spin_week(now, max(weeks) if weeks else None, weeks, members, hist, pot['donors'], settings['elig']) if os.environ.get('SPIN') == '1' else None
    if sp: hist.append(sp); print(f"Quay thưởng {sp['week']}: " + ('Tích trữ' if sp.get('keep') else f"{sp['winner']} trúng {sp['amount']:,.0f}") + f" ({sp['count']} người đủ điều kiện)")
    hist.sort(key=lambda x: (x.get('date', ''), x.get('week', '')))
    pot['history'] = hist
    pot['value'] = sum(x['amount'] for x in pot['donors']) - sum(h.get('amount', 0) for h in hist if not h.get('keep'))
    settings['potBase'] = 0

    # ---- Tàng thư các
    links = []
    ws = need(wb, T_DOCS)
    if ws:
        for r in rows(ws, 2, 4):
            if not nfc(r[1]): continue
            url = nfc(r[2])
            if url and not re.match(r'https?://', url, re.I): url = 'https://' + url
            links.append({'group': nfc(r[0]) or 'Khác', 'title': nfc(r[1]), 'url': url, 'note': nfc(r[3])})

    if ERRORS:
        print('DỪNG LẠI, chưa cập nhật web vì Sheet có lỗi:')
        for e in ERRORS: print('  - ' + e)
        sys.exit(1)

    data = {'version': 5, 'generatedAt': now.strftime('%d/%m/%Y %H:%M'), 'settings': settings, 'notices': notices, 'updates': updates,
            'links': links, 'pot': pot, 'sect': dict(sorted(sect.items())), 'members': members, 'weeks': dict(sorted(weeks.items()))}
    def strip(d): return {k: v for k, v in d.items() if k != 'generatedAt'}
    if old and strip(old) == strip(data):
        print('Không có gì thay đổi.'); return
    with open(args.out, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    act = sum(1 for m in members if not m.get('leftAt'))
    print(f'Đã cập nhật data.json: {act} thành viên, tuần {wk} có {len(entries)} dòng số liệu, {len(weeks)} tuần lưu trữ.')
    for w_ in WARN: print('  Lưu ý: ' + w_)

if __name__ == '__main__':
    main()
